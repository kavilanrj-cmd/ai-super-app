import os
import re
import logging
from datetime import datetime
import pdfplumber
from typing import Optional
from app.agents import agent_coordinator
from app.utils.helpers import extract_skills, calculate_ats_score

logger = logging.getLogger("resume_service")

# Patterns that surface the ATS score the AI actually reported inside its
# Markdown analysis (e.g. "ATS Score: 62/100" or "**Score: 35/100**").
# Ordered from most specific to most generic.
ATS_SCORE_PATTERNS = [
    re.compile(r"(?:ats|overall)\s*score\s*[:=]?\s*(\d{1,3})\s*(?:\/|out\s*of)\s*100", re.IGNORECASE),
    re.compile(r"(?:ats|applicant).{0,100}?(\d{1,3})\s*(?:\/|out\s*of)\s*100", re.IGNORECASE),
    re.compile(r"\b(?:score|candidate)\s*[:=]?\s*(\d{1,3})\s*(?:\/|out\s*of)\s*100", re.IGNORECASE),
    re.compile(r"\b(\d{1,3})\s*(?:\/|out\s*of)\s*100\b"),
]


def extract_ats_score(analysis: str) -> Optional[float]:
    """Return the ATS score the AI reported in its analysis, or None if absent."""
    if not analysis:
        return None
    for pattern in ATS_SCORE_PATTERNS:
        match = pattern.search(analysis)
        if match:
            score = int(match.group(1))
            if 0 <= score <= 100:
                return float(score)
    return None


def _estimate_experience_years(text: str) -> Optional[int]:
    """Best-effort years-of-experience estimate from a resume's text."""
    if not text:
        return None
    patterns = [
        re.compile(r"(\d{1,2})\s*\+?\s*(?:years?|yrs?)\s+of?\s*(?:relevant\s+)?experience", re.IGNORECASE),
        re.compile(r"(\d{1,2})\s*(?:years?|yrs?)\s+experience", re.IGNORECASE),
    ]
    for pat in patterns:
        for m in pat.finditer(text):
            try:
                years = int(m.group(1))
            except (ValueError, IndexError):
                continue
            if 0 < years < 50:
                return years
    # Fall back to counting date ranges like "2018 - 2022" / "2018–2022".
    spans = []
    for m in re.finditer(r"(19|20)\d{2}\s*[-–—/]\s*(present\b|(?:19|20)\d{2})", text, re.IGNORECASE):
        try:
            start = int(re.search(r"(19|20)\d{2}", m.group(0)).group(0))
            end = datetime.now().year if m.group(2).lower() == "present" else int(re.search(r"(19|20)\d{2}", m.group(2)).group(0))
            spans.append(max(0, end - start))
        except Exception:
            continue
    if spans:
        return max(0, min(max(spans), 40))
    return None


def _extract_education(text: str) -> Optional[dict]:
    """Detect the highest education level and school from resume text."""
    if not text:
        return None
    levels = [
        (r"ph\.?d|doctorate", "Doctorate"),
        (r"master(?:'s)?\s+of|m\.?s\.?c|mba|m\.?tech|m\.?e\.?ng", "Master's"),
        (r"bachelor(?:'s)?\s+of|b\.?s\.?c|b\.?tech|b\.?e\.?|b\.?a\.?", "Bachelor's"),
        (r"associate|diploma|hnd", "Diploma / Associate"),
    ]
    best = None
    for pat, label in levels:
        if re.search(pat, text, re.IGNORECASE):
            best = label
            break
    if best is None:
        return None
    school = None
    m = re.search(
        r"(university|institute|college|academy)[^\n.,;|]*",
        text,
        re.IGNORECASE,
    )
    if m:
        school = m.group(0).strip()
    return {"level": best, "school": school}


def _extract_certifications(text: str) -> list:
    """List certification-like mentions (lines containing 'certified' or common cert names)."""
    if not text:
        return []
    cert_keywords = re.compile(
        r"(certified|certification|certificate|pmp|itil|aws\s+certified|scrum|ccna|comptia)",
        re.IGNORECASE,
    )
    found = []
    for line in text.splitlines():
        if cert_keywords.search(line) and len(line.split()) <= 25:
            clean = line.strip("-*• \t")
            if clean and clean not in found:
                found.append(clean[:160])
    return found[:8]


def _extract_job_titles(text: str) -> list:
    """Common job-title lines drawn from the experience section."""
    if not text:
        return []
    title_re = re.compile(
        r"^[ \t]*(?:[•\-*]?\s*)*((?:senior|junior|lead|principal|staff|associate|head\s+of|manager|engineer|developer|designer|analyst|scientist|architect|consultant)[^,\n]{1,60})",
        re.IGNORECASE,
    )
    found = []
    for line in text.splitlines()[:400]:
        if not any(kw in line.lower() for kw in ("engineer", "developer", "designer", "analyst", "manager", "scientist", "architect", "consultant")):
            continue
        m = title_re.match(line)
        if m:
            title = m.group(1).strip().strip("-•* \t")
            if title and title not in found and len(title.split()) <= 8:
                found.append(title)
        if len(found) >= 6:
            break
    return found


def _extract_structured(text: str) -> dict:
    return {
        "experience_years": _estimate_experience_years(text),
        "education": _extract_education(text),
        "certifications": _extract_certifications(text),
        "job_titles": _extract_job_titles(text),
    }

class ResumeService:
    @staticmethod
    async def parse_resume(file_path: str) -> str:
        """Extract text from a resume PDF or DOCX. Never logs the content."""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return await ResumeService._parse_pdf(file_path)
        if ext == ".docx":
            return ResumeService._parse_docx(file_path)
        raise ValueError(f"Unsupported file type: {ext}")

    @staticmethod
    async def _parse_pdf(file_path: str) -> str:
        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text_parts.append(page.extract_text() or "")
        return "\n".join(text_parts)

    @staticmethod
    def _parse_docx(file_path: str) -> str:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs)

    @staticmethod
    async def analyze_resume(file_path: str, job_description: str = "") -> dict:
        try:
            text = await ResumeService.parse_resume(file_path)
        except Exception as e:
            logger.error("Resume parsing failed (file=%s): %s", os.path.basename(file_path), e)
            raise
        if not text.strip():
            logger.warning("Resume produced no extractable text (file=%s)", os.path.basename(file_path))
            raise ValueError("No text could be extracted from the uploaded resume.")

        skills = extract_skills(text)

        analysis = await agent_coordinator.process_with_agent(
            "resume",
            f"Analyze this resume and provide ATS score, missing skills, summary, improvements, and keyword analysis:\n\n{text[:3000]}"
        )

        # The authoritative score is the one the AI computed during analysis.
        # Fall back to the keyword-overlap heuristic only when the AI did not
        # report a score (a genuine 0 is only used when the backend actually
        # returned 0).
        ai_ats_score = extract_ats_score(analysis)
        ats_score = ai_ats_score if ai_ats_score is not None else calculate_ats_score(text, job_description)

        structured = _extract_structured(text)

        return {
            "parsed_text": text,
            "ats_score": ats_score,
            "skills_found": skills,
            "analysis": analysis,
            "experience_years": structured["experience_years"],
            "education": structured["education"],
            "certifications": structured["certifications"],
            "job_titles": structured["job_titles"],
        }
